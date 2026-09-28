import html
import io
import os
import re
import time
import zipfile
from urllib.parse import urljoin

import openpyxl
import pandas as pd
import requests
from dotenv import load_dotenv
from loguru import logger

from utils.database import get_active_workers, get_exclusion_check_notify_users
from utils.google import send_gmail
from utils.misc import json_log_format
from utils.timezone import now_business

logger.add(
    "logs/exclusion-check.log",
    format=json_log_format,
    rotation="50 MB",
    filter=lambda r: r["name"] == "exclusion_check",
)
load_dotenv()

FROM_ADDR = "tech@driftwoodeval.com"

# HHS-OIG List of Excluded Individuals/Entities (LEIE), updated monthly.
OIG_LEIE_URL = "https://oig.hhs.gov/exclusions/downloadables/UPDATED.csv"
OIG_LEIE_REQUIRED_COLUMNS = [
    "LASTNAME",
    "FIRSTNAME",
    "NPI",
    "DOB",
    "EXCLTYPE",
    "EXCLDATE",
]
# LEIE has held around 80k rows for years. A sharp drop means a bad download
# or a header/format change upstream, not a real shrink in the exclusion list.
OIG_LEIE_MIN_ROWS = 50_000

# The download itself is a versioned filename (date + compiler name) that changes
# every update, so the page has to be scraped for the current link each run.
SC_EXCLUSIONS_PAGE_URL = "https://scdhhs.gov/fraud-waste-and-abuse"
# The list has held at 1000+ rows. A sharp drop means a bad download or a
# format change upstream, not a real shrink in the exclusion list.
SC_EXCLUSIONS_MIN_ROWS = 500

# SAM.gov's bulk exclusions extract is free but requires a personal API key
# (Account Details page on sam.gov) and is generated asynchronously: kicking it
# off returns a token URL that returns a "not ready" JSON body until the file
# is built, then the raw CSV.
SAM_EXCLUSIONS_URL = "https://api.sam.gov/entity-information/v4/exclusions"
SAM_DOWNLOAD_URL = "https://api.sam.gov/entity-information/v4/download-exclusions"
# A ~167k-row extract can take a while to build; 8 polls 2 minutes apart gives
# it up to 16 minutes, which is fine for a job that only runs once a month.
SAM_POLL_ATTEMPTS = 8
SAM_POLL_DELAY_SECONDS = 120
# SAM.gov has held around 167k active exclusions. A sharp drop means a bad
# download or a format change upstream, not a real shrink in the exclusion list.
SAM_EXCLUSIONS_MIN_ROWS = 50_000

NAME_SUFFIXES = {"JR", "SR", "II", "III", "IV"}
NAME_TITLES = {"DR", "MR", "MRS", "MS"}


def get_error_recipients() -> list[str]:
    """Parses the ERROR_EMAILS env var, used when there's no DB or nobody has the notifications permission."""
    return [a.strip() for a in os.getenv("ERROR_EMAILS", "").split(",") if a.strip()]


def normalize_name(name: str) -> list[str]:
    """Normalizes a person's name to its first/last name tokens: uppercase, credentials after a comma, titles, suffixes, and punctuation dropped. Middle names/initials are ignored, so key on (first, last)."""
    name = name.split(",", maxsplit=1)[0]
    name = re.sub(r"[^A-Za-z\s]", " ", name).upper()
    words = [w for w in name.split() if w not in NAME_TITLES and w not in NAME_SUFFIXES]
    if len(words) < 2:
        return words
    return [words[0], words[-1]]


def normalize_npi(npi) -> str | None:
    """Returns a 10-digit NPI string, or None if missing/invalid/all-zero."""
    npi = str(npi).strip()
    if not npi.isdigit() or len(npi) > 10:
        return None
    npi = npi.zfill(10)
    return npi if npi != "0000000000" else None


def fetch_oig_leie() -> list[dict]:
    """Downloads and validates the current OIG LEIE exclusion list, returned as exclusion records."""
    response = requests.get(OIG_LEIE_URL, timeout=60)
    response.raise_for_status()

    df = pd.read_csv(io.BytesIO(response.content), dtype=str, keep_default_na=False)

    missing_columns = [c for c in OIG_LEIE_REQUIRED_COLUMNS if c not in df.columns]
    if missing_columns:
        raise ValueError(f"OIG LEIE CSV is missing expected columns: {missing_columns}")

    if len(df) < OIG_LEIE_MIN_ROWS:
        raise ValueError(
            f"OIG LEIE CSV only has {len(df)} rows, expected at least {OIG_LEIE_MIN_ROWS}. "
            "Refusing to treat this as a valid download."
        )

    return [
        {
            "source": "OIG LEIE",
            "npi": normalize_npi(row["NPI"]),
            "first": row["FIRSTNAME"],
            "last": row["LASTNAME"],
            "name": f"{row['FIRSTNAME']} {row['LASTNAME']}".strip(),
            "dob": row["DOB"],
            "city": row["CITY"],
            "state": row["STATE"],
            "exclusion_type": row["EXCLTYPE"],
            "exclusion_date": row["EXCLDATE"],
        }
        for row in df.to_dict("records")
    ]


def find_sc_exclusions_download_url() -> str:
    """Scrapes the SCDHHS excluded providers page for the current .xlsx download link."""
    response = requests.get(SC_EXCLUSIONS_PAGE_URL, timeout=30)
    response.raise_for_status()

    match = re.search(r'href="([^"]+\.xlsx)"', response.text)
    if not match:
        raise ValueError(
            f"Could not find an .xlsx download link on {SC_EXCLUSIONS_PAGE_URL}. "
            "The page layout may have changed."
        )
    return urljoin(response.url, html.unescape(match.group(1)))


def parse_sc_rows(rows: list[tuple]) -> tuple[list[dict], str]:
    """Parses SCDHHS exclusion sheet rows (as returned by openpyxl) into exclusion records.

    Returns the records plus the sheet's own "Reviewed and Updated" banner text,
    which the results email uses as evidence of which list version was screened.
    """
    header_row_index = next(
        (i for i, row in enumerate(rows[:10]) if row and "NPI" in row), None
    )
    if header_row_index is None:
        raise ValueError(
            "Could not find a header row containing 'NPI' in the SCDHHS sheet."
        )

    banner = str(rows[0][0]) if rows and rows[0] and rows[0][0] else ""

    header = [str(c).strip() if c else "" for c in rows[header_row_index]]
    required_columns = [
        "Individual//Entity",
        "NPI",
        "City",
        "State",
        "Excluded/Terminated",
    ]
    missing_columns = [c for c in required_columns if c not in header]
    if missing_columns:
        raise ValueError(
            f"SCDHHS exclusions file is missing expected columns: {missing_columns}"
        )

    col = {name: header.index(name) for name in required_columns}
    date_col = header.index("Action Date") if "Action Date" in header else None

    records = []
    for row in rows[header_row_index + 1 :]:
        entity = str(row[col["Individual//Entity"]] or "").strip()
        if not entity:
            continue

        # Individuals are stored "Last, First[, credentials]"; businesses have no
        # comma and aren't name-matchable against a worker.
        parts = [p.strip() for p in entity.split(",")]
        last = parts[0] if len(parts) > 1 else ""
        first = parts[1] if len(parts) > 1 else ""
        exclusion_date = row[date_col] if date_col is not None else None

        records.append(
            {
                "source": "SC DHHS",
                "npi": normalize_npi(row[col["NPI"]]),
                "first": first,
                "last": last,
                "name": entity,
                "dob": "",
                "city": str(row[col["City"]] or ""),
                "state": str(row[col["State"]] or ""),
                "exclusion_type": str(row[col["Excluded/Terminated"]] or ""),
                "exclusion_date": exclusion_date.strftime("%Y-%m-%d")
                if hasattr(exclusion_date, "strftime")
                else str(exclusion_date or ""),
            }
        )

    return records, banner


def fetch_sc_exclusions() -> tuple[list[dict], str]:
    """Downloads and validates the current SCDHHS excluded/terminated providers list, returned as (exclusion records, list version banner)."""
    download_url = find_sc_exclusions_download_url()
    response = requests.get(download_url, timeout=60)
    response.raise_for_status()

    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))

    records, banner = parse_sc_rows(rows)

    if len(records) < SC_EXCLUSIONS_MIN_ROWS:
        raise ValueError(
            f"SCDHHS exclusions file only has {len(records)} usable rows, expected at least "
            f"{SC_EXCLUSIONS_MIN_ROWS}. Refusing to treat this as a valid download."
        )

    return records, banner


# Confirmed against GSA's published sample extract (open.gsa.gov/api/exclusions-api/
# v4/exclusion-sample-csv-1.xlsx). The schema has no exclusion-date field.
SAM_REQUIRED_COLUMNS = [
    "classificationType",
    "npi",
    "firstName",
    "lastName",
    "entityName",
    "exclusionType",
    "exclusionProgram",
    "exclusionPrimaryAddress:city",
    "exclusionPrimaryAddress:stateOrProvinceCode",
]


def _redact(text: str, secret: str) -> str:
    """Replaces every occurrence of secret in text, so an API key never reaches the log."""
    return text.replace(secret, "***") if secret else text


def fetch_sam_exclusions() -> list[dict]:
    """Downloads and validates the current SAM.gov active exclusions extract, returned as exclusion records.

    Requires SAM_GOV_API_KEY (a free personal API key from the SAM.gov Account
    Details page). The extract is generated asynchronously: this polls the
    download URL until the file is ready or the poll budget runs out.
    """
    api_key = os.getenv("SAM_GOV_API_KEY")
    if not api_key:
        raise ValueError(
            "SAM_GOV_API_KEY is not set. Generate a free personal API key from the "
            "Account Details page on sam.gov and set it in .env."
        )

    try:
        kickoff = requests.get(
            SAM_EXCLUSIONS_URL,
            params={"api_key": api_key, "format": "csv", "recordStatus": "active"},
            timeout=30,
        )
        if kickoff.status_code == 429:
            raise ValueError("SAM.gov daily request limit reached, try again tomorrow.")
        kickoff.raise_for_status()
    except requests.RequestException as e:
        raise RuntimeError(_redact(str(e), api_key)) from None

    # The message embeds a literal "REPLACE_WITH_API_KEY" placeholder alongside the
    # real token, so the token is pulled out and the download request is built
    # directly rather than reusing the URL as-is.
    message = kickoff.json().get("message", "")
    token_match = re.search(r"token=([^&\s]+)", message)
    if not token_match:
        raise ValueError(
            f"Unexpected SAM.gov extract kickoff response: {_redact(message, api_key)}"
        )
    token = token_match.group(1)

    csv_bytes = None
    for attempt in range(1, SAM_POLL_ATTEMPTS + 1):
        time.sleep(SAM_POLL_DELAY_SECONDS)
        try:
            poll = requests.get(
                SAM_DOWNLOAD_URL,
                params={"api_key": api_key, "token": token},
                timeout=60,
            )
            poll.raise_for_status()
        except requests.RequestException as e:
            raise RuntimeError(_redact(str(e), api_key)) from None

        if "json" in poll.headers.get("Content-Type", ""):
            logger.info(
                f"SAM.gov extract not ready yet (attempt {attempt}/{SAM_POLL_ATTEMPTS})"
            )
            continue
        csv_bytes = poll.content
        break

    if csv_bytes is None:
        raise TimeoutError(
            f"SAM.gov exclusions extract did not become ready after {SAM_POLL_ATTEMPTS} polls."
        )

    if csv_bytes[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(csv_bytes)) as archive:
            csv_name = next(n for n in archive.namelist() if n.lower().endswith(".csv"))
            csv_bytes = archive.read(csv_name)

    df = pd.read_csv(io.BytesIO(csv_bytes), dtype=str, keep_default_na=False)

    missing_columns = [c for c in SAM_REQUIRED_COLUMNS if c not in df.columns]
    if missing_columns:
        raise ValueError(
            f"SAM.gov exclusions CSV is missing expected columns: {missing_columns}. "
            f"Got columns: {list(df.columns)}"
        )

    if len(df) < SAM_EXCLUSIONS_MIN_ROWS:
        raise ValueError(
            f"SAM.gov exclusions CSV only has {len(df)} rows, expected at least "
            f"{SAM_EXCLUSIONS_MIN_ROWS}. Refusing to treat this as a valid download."
        )

    records = [
        {
            "source": "SAM.gov",
            "npi": normalize_npi(row["npi"]) if row["npi"] != "null" else None,
            "first": "" if row["firstName"] == "null" else row["firstName"],
            "last": "" if row["lastName"] == "null" else row["lastName"],
            "name": row["entityName"]
            if row["entityName"] != "null"
            else f"{row['firstName']} {row['lastName']}".strip(),
            "dob": "",
            "city": ""
            if row["exclusionPrimaryAddress:city"] == "null"
            else row["exclusionPrimaryAddress:city"],
            "state": ""
            if row["exclusionPrimaryAddress:stateOrProvinceCode"] == "null"
            else row["exclusionPrimaryAddress:stateOrProvinceCode"],
            "exclusion_type": row["exclusionType"],
            "exclusion_date": "",
        }
        for row in df.to_dict("records")
    ]

    name_screenable = sum(
        1 for r in records if len(normalize_name(f"{r['first']} {r['last']}")) == 2
    )
    if records and name_screenable == 0:
        raise ValueError(
            "No SAM.gov record produced a name-matchable (first, last) pair. "
            "The firstName/lastName columns may have changed shape."
        )

    return records


def dedupe_matches(matches: list[dict]) -> list[dict]:
    """Collapses duplicate matches from a worker who appears as both a user and an evaluator row."""
    seen = set()
    deduped = []
    for match in matches:
        key = (
            match["worker_email"],
            match["source"],
            match["match_name"],
            match["match_npi"],
        )
        if key not in seen:
            seen.add(key)
            deduped.append(match)
    return deduped


def match_workers(workers: list[dict], exclusion_records: list[dict]) -> list[dict]:
    """Screens workers against a combined list of exclusion records, matching by NPI and by normalized (first, last) name. Returns one dict per match with worker and record details.

    lazy: name matches for a common name re-alert every month with no way to mark
    them reviewed and cleared. Upgrade path: an ignore table keyed on worker email
    plus record source/NPI/name/DOB, filled from a settings page, checked here
    before a match is included.
    """
    npi_index: dict[str, int] = {}
    name_index: dict[tuple[str, str], list[int]] = {}

    for i, record in enumerate(exclusion_records):
        if record["npi"]:
            npi_index[record["npi"]] = i

        if record["last"]:
            key = normalize_name(f"{record['first']} {record['last']}")
            if len(key) == 2:
                name_index.setdefault((key[0], key[1]), []).append(i)

    matches = []
    for worker in workers:
        hit_indices: set[int] = set()

        worker_npi = normalize_npi(worker["npi"]) if worker["npi"] else None
        if worker_npi:
            npi_hit = npi_index.get(worker_npi)
            if npi_hit is not None:
                hit_indices.add(npi_hit)

        if worker["name"]:
            key = normalize_name(worker["name"])
            if len(key) == 2:
                hit_indices.update(name_index.get((key[0], key[1]), []))

        matches.extend(
            {
                "worker_name": worker["name"],
                "worker_email": worker["email"],
                "source": exclusion_records[i]["source"],
                "match_name": exclusion_records[i]["name"],
                "match_npi": exclusion_records[i]["npi"] or "",
                "match_dob": exclusion_records[i]["dob"],
                "match_city": exclusion_records[i]["city"],
                "match_state": exclusion_records[i]["state"],
                "match_type": exclusion_records[i]["exclusion_type"],
                "match_date": exclusion_records[i]["exclusion_date"],
            }
            for i in hit_indices
        )

    return dedupe_matches(matches)


def send_results_email(
    matches: list[dict],
    worker_count: int,
    unscreened_emails: list[str],
    source_counts: dict[str, int],
    failed_sources: list[str],
) -> None:
    """Emails the monthly screening result to users with the notifications permission, falling back to ERROR_EMAILS if nobody has it yet.

    With DEV_TOGGLE set, always sends to ERROR_EMAILS only, so a manual run
    doesn't email real staff about real workers.
    """
    if os.getenv("DEV_TOGGLE"):
        recipients = get_error_recipients()
        logger.info("Dev mode: sending exclusion check results to ERROR_EMAILS only.")
    else:
        recipients = [u["email"] for u in get_exclusion_check_notify_users()]
        if not recipients:
            recipients = get_error_recipients()
            logger.warning(
                "No users have settings:exclusion-check:notifications, falling back to ERROR_EMAILS."
            )
    if not recipients:
        logger.error(
            "No settings:exclusion-check:notifications recipients and ERROR_EMAILS is unset, "
            "exclusion check results were not sent anywhere."
        )
        return

    run_date = now_business().strftime("%Y-%m-%d")
    subject = f"Exclusion List Check ({run_date}): {len(matches)} match(es)"

    sources_summary = ", ".join(
        f"{count} {source}" for source, count in source_counts.items()
    )
    summary = (
        f"Screened {worker_count} workers against {sources_summary} on {run_date}."
    )
    if unscreened_emails:
        summary += f" {len(unscreened_emails)} worker(s) have no usable first/last name on file and were not name-screened: {', '.join(unscreened_emails)}."
    if failed_sources:
        summary += f" WARNING: {', '.join(failed_sources)} could not be checked this run, see the log."

    if matches:
        rows_html = "".join(
            f"<tr><td>{html.escape(str(m['worker_name']))}</td><td>{html.escape(m['worker_email'])}</td>"
            f"<td>{html.escape(m['source'])}</td><td>{html.escape(m['match_name'])}</td><td>{html.escape(m['match_npi'])}</td>"
            f"<td>{html.escape(m['match_dob'])}</td><td>{html.escape(m['match_city'])}, {html.escape(m['match_state'])}</td>"
            f"<td>{html.escape(m['match_type'])}</td><td>{html.escape(str(m['match_date']))}</td></tr>"
            for m in matches
        )
        html_content = f"""
        <p>{html.escape(summary)}
        <strong>{len(matches)} possible match(es) found.</strong> Review each one manually,
        name matches alone are not conclusive.</p>
        <table border="1" cellpadding="4" cellspacing="0">
            <tr><th>Worker</th><th>Email</th><th>Source</th><th>Matched Name</th><th>NPI</th>
                <th>DOB</th><th>Location</th><th>Excl. Type</th><th>Excl. Date</th></tr>
            {rows_html}
        </table>
        """
        message_text = f"{summary} {len(matches)} possible match(es) found. See the HTML version for details."
    else:
        html_content = f"<p>{html.escape(summary)} No matches found.</p>"
        message_text = f"{summary} No matches found."

    for recipient in recipients:
        send_gmail(
            message_text=message_text,
            subject=subject,
            to_addr=recipient,
            from_addr=FROM_ADDR,
            html=html_content,
        )
    logger.info(
        f"Sent exclusion check results to {len(recipients)} recipient(s), {len(matches)} match(es)."
    )


def send_failure_email(error: Exception) -> None:
    """Emails ERROR_EMAILS directly, bypassing the DB, since a failed run is itself a compliance gap."""
    recipients = get_error_recipients()
    if not recipients:
        logger.error(
            "ERROR_EMAILS is not set, cannot send exclusion check failure notice."
        )
        return

    run_date = now_business().strftime("%Y-%m-%d")
    send_gmail(
        message_text=f"The monthly exclusion list check failed on {run_date}: {error}\n\nNo workers were screened this run. Check the exclusion-check log.",
        subject=f"Exclusion List Check FAILED ({run_date})",
        to_addr=", ".join(recipients),
        from_addr=FROM_ADDR,
    )


def fetch_all_exclusion_records() -> tuple[list[dict], dict[str, int], list[str]]:
    """Fetches every exclusion source independently, so one source failing doesn't stop the others.

    Returns the combined records, a count of records per successful source, and
    the names of any sources that failed.
    """
    records: list[dict] = []
    source_counts: dict[str, int] = {}
    failed_sources: list[str] = []

    logger.info("Fetching OIG LEIE exclusions")
    try:
        oig_records = fetch_oig_leie()
        records.extend(oig_records)
        source_counts["OIG LEIE"] = len(oig_records)
        logger.info(f"Fetched {len(oig_records)} OIG LEIE records")
    except Exception:
        logger.exception("Failed to fetch OIG LEIE")
        failed_sources.append("OIG LEIE")

    logger.info("Fetching SC DHHS exclusions")
    try:
        sc_records, sc_banner = fetch_sc_exclusions()
        records.extend(sc_records)
        source_counts[f"SC DHHS ({sc_banner or 'version unknown'})"] = len(sc_records)
        logger.info(
            f"Fetched {len(sc_records)} SC DHHS records ({sc_banner or 'version unknown'})"
        )
    except Exception:
        logger.exception("Failed to fetch SC DHHS exclusions")
        failed_sources.append("SC DHHS")

    logger.info("Fetching SAM.gov exclusions")
    try:
        sam_records = fetch_sam_exclusions()
        records.extend(sam_records)
        source_counts["SAM.gov"] = len(sam_records)
        logger.info(f"Fetched {len(sam_records)} SAM.gov records")
    except Exception:
        logger.exception("Failed to fetch SAM.gov exclusions")
        failed_sources.append("SAM.gov")

    if not source_counts:
        raise RuntimeError(f"All exclusion sources failed: {', '.join(failed_sources)}")

    return records, source_counts, failed_sources


def main():
    logger.info("Starting monthly exclusion check")
    try:
        exclusion_records, source_counts, failed_sources = fetch_all_exclusion_records()
        logger.info(
            f"Screening against {len(exclusion_records)} exclusion record(s) "
            f"from {len(source_counts)} source(s), {len(failed_sources)} source(s) failed"
        )

        workers = get_active_workers()
        worker_count = len({w["email"].lower() for w in workers})

        screened_emails = {
            w["email"].lower()
            for w in workers
            if len(normalize_name(w["name"] or "")) == 2
        }
        unscreened_emails = sorted(
            {w["email"] for w in workers if w["email"].lower() not in screened_emails}
        )

        matches = match_workers(workers, exclusion_records)
        logger.info(
            f"Screened {worker_count} worker(s), found {len(matches)} match(es)"
        )
        send_results_email(
            matches,
            worker_count=worker_count,
            unscreened_emails=unscreened_emails,
            source_counts=source_counts,
            failed_sources=failed_sources,
        )
    except Exception as e:
        logger.exception(f"Failed to run exclusion check: {e}")
        try:
            send_failure_email(e)
        except Exception:
            logger.exception("Failed to send exclusion check failure notice")


if __name__ == "__main__":
    main()
